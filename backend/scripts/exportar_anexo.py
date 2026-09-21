"""Genera el anexo de la base de datos consolidada, en Excel.

Es el medio de verificación de la actividad 3.2: la base de datos del proyecto
en un archivo que se puede abrir sin programar, con su diccionario de variables
y el estado de cada una.

    python backend/scripts/exportar_anexo.py

Lee el CSV que deja `cargar_plantilla.py`, así que el orden es siempre:
primero cargar la plantilla, después generar el anexo. `actualizar.py` hace
las dos cosas seguidas.

Hojas del archivo:

    Léame         Qué es el archivo y cómo leerlo.
    Datos         Un renglón por lote, tal como quedó consolidado.
    Diccionario   Las diecisiete variables (V01 a V17): qué mide cada una,
                  en qué unidad, para qué la usa el modelo y de dónde sale.
    Completitud   Cuántos lotes tienen cada variable. Deja a la vista lo que
                  falta y por qué, en lugar de esconderlo.
    Dietas        El catálogo de las tres comidas en estudio.
"""

from __future__ import annotations

import argparse
import csv
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Diccionario de variables
# --------------------------------------------------------------------------
# Los títulos y unidades son los de la plantilla de recolección, para que el
# anexo y la plantilla hablen el mismo idioma. Cada entrada:
#     (código, columna del CSV, nombre, unidad, papel en el modelo, origen)
#
# El origen importa: dice QUIÉN produce el dato. Lo que depende del
# laboratorio no puede estar completo mientras el laboratorio no entregue, y
# eso no es una falla de la base de datos.

LABORATORIO = "Análisis bromatológico (laboratorio)"
PLANTILLA = "Plantilla de recolección"
CIERRE = "Plantilla, al cierre del ciclo"
PLATAFORMA = "Consulta en la plataforma"

DICCIONARIO = [
    ("V01", "tipo_dieta", "Código de dieta", "código", "Entrada del modelo", PLANTILLA),
    ("V02", "alimento_g_dia", "Alimento promedio", "g/día", "Entrada del modelo", PLANTILLA),
    ("V03", "temperatura", "Temperatura promedio", "°C", "Entrada del modelo", PLANTILLA),
    ("V04", "humedad_ambiental", "Humedad promedio", "%", "Entrada del modelo", PLANTILLA),
    ("V05", "fotoperiodo", "Fotoperiodo", "h luz/día", "Registro de cría", PLANTILLA),
    ("V06", "densidad", "Densidad inicial", "grillos/m²", "Registro de cría", PLANTILLA),
    ("V07", "especie", "Grupo de grillo", "categoría", "Entrada del modelo", PLANTILLA),
    ("V08", "tiempo_desarrollo", "Tiempo de desarrollo", "días", "Entrada del modelo", PLANTILLA),
    ("V09", "n_grillos_inicio", "Grillos al inicio", "individuos", "Registro de cría", PLANTILLA),
    ("V10", "proteina_harina", "Proteína en harina", "% peso seco", "Variable objetivo", LABORATORIO),
    ("V11", "lipidos_harina", "Lípidos en harina", "% peso seco", "Variable objetivo", LABORATORIO),
    ("V12", "tasa_supervivencia", "Supervivencia", "%", "Filtro de viabilidad", PLANTILLA),
    ("V13", "longitud_final", "Longitud final promedio", "mm", "Filtro de viabilidad", CIERRE),
    ("V14", "biomasa_total", "Biomasa total", "g", "Filtro de viabilidad", CIERRE),
    ("V15", None, "Especie destino", "categoría", "Contexto de la consulta", PLATAFORMA),
    ("V16", None, "Etapa de desarrollo del animal", "categoría", "Contexto de la consulta", PLATAFORMA),
    ("V17", None, "Requerimiento proteico del destino", "%", "Contexto de la consulta", PLATAFORMA),
]

#: Columnas del CSV que no son variables del catálogo pero sí van en los datos.
COLUMNAS_EXTRA = [
    ("id_ensayo", "Identificador del lote"),
    ("fuente", "Archivo de origen"),
    ("observaciones", "Observaciones"),
]

DIETAS = [
    ("D1", "Harina de bore (80%), harina de choclo (10%), avena en hojuelas (10%)"),
    ("D2", "Harina de botón de oro (80%), harina de choclo (10%), avena en hojuelas (10%)"),
    ("D3", "Salvado de trigo (80%), harina de choclo (10%), avena en hojuelas (10%)"),
]

# --------------------------------------------------------------------------
# Limpieza para un documento que sale del proyecto
# --------------------------------------------------------------------------

#: La plantilla anota el nombre de especie que se usó al registrar el lote.
#: La identificación a nivel de especie NO está confirmada, así que un anexo
#: que se entrega fuera del proyecto no puede afirmarla. Se quita esa frase;
#: el dato sigue intacto en la plantilla original.
ESPECIE_EN_NOTAS = re.compile(r"Especie según plantilla:[^.]*\.\s*", re.IGNORECASE)

#: Cómo se presenta V07 mientras no haya identificación de especie.
GRUPO_DE_GRILLO = "Familia Gryllidae (especie por confirmar)"


def limpiar(fila: dict) -> dict:
    """Prepara un lote para el anexo sin tocar el dato de origen."""
    salida = dict(fila)
    salida["observaciones"] = ESPECIE_EN_NOTAS.sub("", fila.get("observaciones") or "").strip()
    if salida.get("especie"):
        salida["especie"] = GRUPO_DE_GRILLO
    # La fuente trae el nombre del Excel con "(1)" y demás restos de
    # descargas; lo que interesa es de dónde vino, no el nombre del archivo.
    fuente = (fila.get("fuente") or "").split("·")[0].strip()
    salida["fuente"] = fuente.capitalize() if fuente else ""
    return salida


def numero(valor: str):
    """El CSV trae todo como texto; en Excel los números deben ser números."""
    if valor in (None, ""):
        return None
    try:
        n = float(valor)
    except ValueError:
        return valor
    return int(n) if n.is_integer() else round(n, 4)


# --------------------------------------------------------------------------
# Estilo
# --------------------------------------------------------------------------

VERDE = "2F4A2C"
ENCABEZADO = PatternFill("solid", fgColor=VERDE)
PENDIENTE = PatternFill("solid", fgColor="FFF4D6")  # ámbar suave
NO_APLICA = PatternFill("solid", fgColor="EEEEEE")
BLANCO = Font(color="FFFFFF", bold=True)
NEGRITA = Font(bold=True)
AJUSTE = Alignment(wrap_text=True, vertical="top")


def encabezar(hoja, titulos: list[str]) -> None:
    hoja.append(titulos)
    for celda in hoja[1]:
        celda.fill = ENCABEZADO
        celda.font = BLANCO
        celda.alignment = AJUSTE
    hoja.freeze_panes = "A2"


def anchos(hoja, medidas: list[int]) -> None:
    for i, ancho in enumerate(medidas, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


# --------------------------------------------------------------------------
# Hojas
# --------------------------------------------------------------------------


def hoja_leame(libro: Workbook, lotes: list[dict], completitud: list[tuple]) -> None:
    h = libro.active
    h.title = "Léame"
    pendientes = [c for c in completitud if c[5] == "Pendiente de laboratorio"]
    lineas = [
        ("Base de datos consolidada · GrillosIA", NEGRITA),
        ("Proyecto Cría de grillos y optimización de su dieta mediante inteligencia artificial", None),
        ("Convocatoria 963-2025 Minciencias · Contrato 207-2025 · Universidad de los Llanos", None),
        ("", None),
        (f"Generado el {date.today().isoformat()} a partir de la plantilla de recolección.", None),
        (f"Lotes consolidados: {len(lotes)}.", None),
        ("", None),
        ("Qué contiene", NEGRITA),
        ("Datos: un renglón por lote, con toda la información registrada hasta la fecha de generación.", None),
        ("Diccionario: las diecisiete variables del proyecto (V01 a V17), su unidad, su papel en el modelo y quién produce el dato.", None),
        ("Completitud: cuántos lotes tienen cada variable, y por qué falta lo que falta.", None),
        ("Dietas: el catálogo de las tres comidas en estudio.", None),
        ("", None),
        ("Cómo leer los vacíos", NEGRITA),
        (
            "Las celdas en ámbar corresponden a proteína (V10) y lípidos (V11) en harina. "
            "Se obtienen por análisis bromatológico al cierre del ciclo productivo y se "
            "incorporan a medida que el laboratorio entrega resultados. Que estén vacías "
            "no es un error de registro.",
            None,
        ),
        (
            "Las variables V15 a V17 describen al animal que va a consumir la harina. "
            "No se miden en la cría: las aporta el productor al hacer su consulta en la plataforma.",
            None,
        ),
        (
            "El grupo de grillo figura como Familia Gryllidae: la identificación a nivel de "
            "especie está pendiente de confirmación.",
            None,
        ),
    ]
    if pendientes:
        lineas.append(("", None))
        lineas.append(
            (
                f"Variables a la espera de laboratorio: {', '.join(p[0] for p in pendientes)}.",
                None,
            )
        )
    for texto, fuente in lineas:
        h.append([texto])
        celda = h.cell(row=h.max_row, column=1)
        celda.alignment = AJUSTE
        if fuente:
            celda.font = fuente
    h.column_dimensions["A"].width = 110


def hoja_datos(libro: Workbook, lotes: list[dict]) -> None:
    h = libro.create_sheet("Datos")
    columnas = [("id_ensayo", "Identificador del lote")]
    columnas += [(col, f"{cod} · {nombre} ({unidad})") for cod, col, nombre, unidad, *_ in DICCIONARIO if col]
    columnas += [("fuente", "Fuente"), ("observaciones", "Observaciones")]

    encabezar(h, [titulo for _, titulo in columnas])
    objetivos = {"proteina_harina", "lipidos_harina"}
    for lote in lotes:
        h.append([numero(lote.get(col)) if col not in ("id_ensayo", "tipo_dieta", "especie", "fuente", "observaciones") else lote.get(col) for col, _ in columnas])
        fila = h.max_row
        for i, (col, _) in enumerate(columnas, start=1):
            celda = h.cell(row=fila, column=i)
            celda.alignment = AJUSTE
            if col in objetivos and celda.value in (None, ""):
                celda.fill = PENDIENTE
    anchos(h, [14] + [16] * (len(columnas) - 3) + [22, 60])


def hoja_diccionario(libro: Workbook) -> None:
    h = libro.create_sheet("Diccionario")
    encabezar(h, ["Código", "Variable", "Unidad", "Columna en la base", "Papel en el modelo", "Origen del dato"])
    for cod, col, nombre, unidad, papel, origen in DICCIONARIO:
        h.append([cod, nombre, unidad, col or "(no se almacena en la cría)", papel, origen])
        for celda in h[h.max_row]:
            celda.alignment = AJUSTE
    anchos(h, [9, 34, 14, 26, 24, 36])


def calcular_completitud(lotes: list[dict]) -> list[tuple]:
    """Una fila por variable: código, nombre, con dato, total, porcentaje, estado."""
    total = len(lotes)
    filas = []
    for cod, col, nombre, _, _, origen in DICCIONARIO:
        if col is None:
            filas.append((cod, nombre, None, None, None, "No aplica a la cría"))
            continue
        con_dato = sum(1 for lote in lotes if (lote.get(col) or "") != "")
        if con_dato == total:
            estado = "Completa"
        elif origen == LABORATORIO:
            estado = "Pendiente de laboratorio"
        elif origen == CIERRE:
            # Longitud final y biomasa se miden al cosechar. Un lote que sigue
            # en cría no puede tenerlas todavía; uno ya cosechado sí debería.
            estado = "Se registra al cierre del ciclo"
        elif con_dato == 0:
            estado = "Sin registrar"
        else:
            estado = "Parcial"
        filas.append((cod, nombre, con_dato, total, con_dato / total if total else 0, estado))
    return filas


def hoja_completitud(libro: Workbook, completitud: list[tuple]) -> None:
    h = libro.create_sheet("Completitud")
    encabezar(h, ["Código", "Variable", "Lotes con dato", "Lotes totales", "Completitud", "Estado"])
    for fila in completitud:
        h.append(list(fila))
        ultima = h.max_row
        h.cell(row=ultima, column=5).number_format = "0%"
        estado = fila[5]
        relleno = PENDIENTE if estado == "Pendiente de laboratorio" else NO_APLICA if estado == "No aplica a la cría" else None
        if relleno:
            for celda in h[ultima]:
                celda.fill = relleno
    anchos(h, [9, 34, 15, 14, 13, 26])


def hoja_dietas(libro: Workbook) -> None:
    h = libro.create_sheet("Dietas")
    encabezar(h, ["Código", "Composición", "Hidratación"])
    for cod, composicion in DIETAS:
        h.append([cod, composicion, "Manzana"])
        for celda in h[h.max_row]:
            celda.alignment = AJUSTE
    anchos(h, [9, 80, 16])


# --------------------------------------------------------------------------


def generar(csv_entrada: Path, salida: Path) -> list[tuple]:
    with csv_entrada.open(encoding="utf-8") as f:
        lotes = [limpiar(fila) for fila in csv.DictReader(f)]
    if not lotes:
        raise SystemExit(f"{csv_entrada} no tiene lotes.")

    completitud = calcular_completitud(lotes)
    libro = Workbook()
    hoja_leame(libro, lotes, completitud)
    hoja_datos(libro, lotes)
    hoja_diccionario(libro)
    hoja_completitud(libro, completitud)
    hoja_dietas(libro)

    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    return completitud


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--csv", default="data/experimentos.csv", help="CSV consolidado de entrada")
    p.add_argument("--salida", default="entregables/base-datos-consolidada.xlsx", help="Archivo de Excel de salida")
    args = p.parse_args()

    completitud = generar(Path(args.csv), Path(args.salida))
    print(f"Anexo generado: {args.salida}\n")
    print(f"  {'Código':<7} {'Variable':<36} {'Con dato':>9}  Estado")
    for cod, nombre, con_dato, total, _, estado in completitud:
        cuenta = f"{con_dato}/{total}" if total else "—"
        print(f"  {cod:<7} {nombre:<36} {cuenta:>9}  {estado}")


if __name__ == "__main__":
    main()
